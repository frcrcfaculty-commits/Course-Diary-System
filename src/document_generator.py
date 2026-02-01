"""
Course Diary Automation System
Document Generator for creating Word documents from course data
"""

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
import os
from datetime import datetime


class CourseDiaryGenerator:
    """Generate Course Diary Word documents from JSON data"""
    
    def __init__(self, global_config_path: str):
        with open(global_config_path, 'r') as f:
            self.global_config = json.load(f)
        self.doc = None
    
    def create_document(self, course_data: dict, output_path: str):
        """Create a complete course diary document"""
        self.doc = Document()
        self._setup_styles()
        self._add_header(course_data)  # Add header to all pages
        
        # Generate all sections
        self._add_title_page(course_data)
        self._add_mission_vision()  # NEW: Add Mission/Vision page
        self._add_document_index(course_data)
        self._add_personal_timetable(course_data)
        self._add_course_outcomes(course_data)
        self._add_lesson_plan(course_data)
        self._add_co_po_mapping(course_data)
        self._add_assessment_tools(course_data)
        self._add_recommended_books(course_data)
        self._add_exit_survey(course_data)
        
        # Save document
        self.doc.save(output_path)
        print(f"Document saved to: {output_path}")
        return output_path
    
    def _add_header(self, course_data: dict):
        """Add header with 'Fr. C.R.C.E. Course Diary | Academic Year' to all pages"""
        section = self.doc.sections[0]
        header = section.header
        
        # Create a table in the header
        header_table = header.add_table(rows=1, cols=2, width=Inches(6.5))
        
        # Left cell: "Fr. C.R.C.E.   Course Diary"
        left_cell = header_table.rows[0].cells[0]
        left_para = left_cell.paragraphs[0]
        left_run = left_para.add_run("Fr. C.R.C.E.")
        left_run.font.size = Pt(10)
        left_para.add_run("\t\t\t\t\t\t      Course Diary")
        
        # Right cell: Academic Year
        right_cell = header_table.rows[0].cells[1]
        right_para = right_cell.paragraphs[0]
        right_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        academic_year = course_data.get('course_info', {}).get('academic_year', '2025-26')
        year_run = right_para.add_run(academic_year)
        year_run.font.size = Pt(10)
    
    def _add_mission_vision(self):
        """Add Mission/Vision statement page"""
        info = self.global_config['institution']
        
        # VISION section
        p = self.doc.add_paragraph()
        run = p.add_run("VISION:")
        run.bold = True
        run.font.size = Pt(12)
        
        p = self.doc.add_paragraph()
        p.add_run(info['vision'])
        
        self.doc.add_paragraph()
        
        # MISSION section
        p = self.doc.add_paragraph()
        run = p.add_run("MISSION:")
        run.bold = True
        run.font.size = Pt(12)
        
        for mission in info['mission']:
            p = self.doc.add_paragraph()
            p.add_run(f"• {mission}")
        
        self.doc.add_paragraph()
        
        # PEOs section
        p = self.doc.add_paragraph()
        run = p.add_run("PROGRAM EDUCATIONAL OBJECTIVES (PEOs):")
        run.bold = True
        run.font.size = Pt(12)
        
        p = self.doc.add_paragraph()
        p.add_run("The graduates of BE in Artificial Intelligence & Data Science will be able to:")
        
        for peo in info['peos']:
            p = self.doc.add_paragraph()
            p.add_run(f"• {peo}")
        
        self.doc.add_paragraph()
        
        # PSO section
        p = self.doc.add_paragraph()
        run = p.add_run("PROGRAM SPECIFIC OUTCOME (PSO):")
        run.bold = True
        run.font.size = Pt(12)
        
        p = self.doc.add_paragraph()
        p.add_run(info['pso'])
        
        self.doc.add_page_break()
    
    def _setup_styles(self):
        """Configure document styles"""
        # Set default font
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Times New Roman'
        font.size = Pt(12)
        
        # Heading styles
        for i in range(1, 4):
            heading_style = self.doc.styles[f'Heading {i}']
            heading_style.font.name = 'Times New Roman'
            heading_style.font.bold = True
    
    def _add_title_page(self, course_data: dict):
        """Add title page with course info"""
        info = self.global_config['institution']
        course = course_data['course_info']
        
        # Department header
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Department of {info['department']}")
        run.bold = True
        run.font.size = Pt(16)
        
        self.doc.add_paragraph()
        
        # Course File Index title
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Course File Index")
        run.bold = True
        run.font.size = Pt(18)
        
        # Academic Year
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(course['academic_year'])
        run.font.size = Pt(14)
        
        self.doc.add_paragraph()
        
        # Course details
        p = self.doc.add_paragraph()
        p.add_run(f"Course Name: ").bold = True
        p.add_run(course['course_name'])
        
        p = self.doc.add_paragraph()
        p.add_run(f"Course ID: ").bold = True
        p.add_run(course['course_code'])
        p.add_run("    ")
        p.add_run(f"Semester: ").bold = True
        p.add_run(course['semester'])
        
        self.doc.add_page_break()
    
    def _add_document_index(self, course_data: dict):
        """Add document index table matching the official format"""
        self.doc.add_heading('Document Index', level=1)
        
        course_type = course_data['course_info']['course_type']
        
        # Define documents based on course type - Updated to match NLP diary format
        if course_type == 'lab':
            docs = [
                "Department Mission and Vision",
                "Personal timetable, Syllabus",
                "Lab Plan",
                "Lab Outcomes, PO mapping with justification",
                "Lab Manual",
                "ISE1 and ISE2 (Lab) evaluation tools and documents",
                "ISE Grading sheets / Attendance sheet",
                "List of Mini Projects with details",
                "Lab Exit Survey",
                "LO and PO Attainments",
            ]
        else:  # theory or theory_tutorial
            docs = [
                "Department Mission and Vision",
                "Personal timetable, Syllabus",
                "Course & Program Outcomes (CO & PO) Mapping, Justification & lesson plan",
                "MSE paper with Solution",
                "ESE paper with Solution",
                "Special exam paper with Solution",
                "ISE1 and ISE2 (Theory) evaluation tools and documents, attendance",
                "Beyond syllabus activity/ Innovative method used / Project based learning report / Guest Lecture (proof)",
                "Identification of Slow and Fast Learners",
                "Course Exit Survey",
                "ISE / MSE / ESE Grading sheets",
                "MSE and ESE Analysis Report",
                "CO and PO Attainments",
            ]
        
        # Create table
        table = self.doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        
        # Header row
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Sr. No.'
        hdr_cells[1].text = 'List of Documents Attached'
        hdr_cells[2].text = 'Tag No.'
        
        for cell in hdr_cells:
            cell.paragraphs[0].runs[0].bold = True
        
        # Data rows
        for i, description in enumerate(docs, 1):
            row_cells = table.add_row().cells
            row_cells[0].text = str(i)
            row_cells[1].text = description
            row_cells[2].text = ''  # Tag No. left empty for user to fill
        
        self.doc.add_page_break()
    
    def _add_personal_timetable(self, course_data: dict):
        """Add personal timetable section"""
        self.doc.add_heading('Personal Time Table', level=1)
        
        # Faculty name
        p = self.doc.add_paragraph()
        p.add_run("Name of the Faculty: ").bold = True
        p.add_run(course_data['course_info']['faculty_name'])
        
        # Create timetable
        time_slots = self.global_config['time_slots']
        days = self.global_config['days']
        
        table = self.doc.add_table(rows=len(days) + 1, cols=len(time_slots) + 1)
        table.style = 'Table Grid'
        
        # Header row
        table.rows[0].cells[0].text = 'Day'
        for i, slot in enumerate(time_slots):
            table.rows[0].cells[i + 1].text = slot
            table.rows[0].cells[i + 1].paragraphs[0].runs[0].font.size = Pt(8)
        
        # Day rows
        for i, day in enumerate(days):
            table.rows[i + 1].cells[0].text = day[:3]  # Mon, Tue, etc.
        
        self.doc.add_page_break()
    
    def _add_course_outcomes(self, course_data: dict):
        """Add Course Outcomes section"""
        info = course_data['course_info']
        
        # Header
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(info['branch'])
        run.bold = True
        
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(f"(Academic Year: {info['academic_year']})")
        
        # Course details
        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.add_run("Course Code: ").bold = True
        p.add_run(info['course_code'])
        
        p = self.doc.add_paragraph()
        p.add_run("Course Name: ").bold = True
        p.add_run(info['course_name'])
        
        p = self.doc.add_paragraph()
        p.add_run("Course Teacher: ").bold = True
        p.add_run(info['faculty_name'])
        
        # Course Outcomes
        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.add_run("Course Outcomes (CO): ").bold = True
        p.add_run("At the End of the course students will be able to")
        
        # CO table
        table = self.doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        
        hdr = table.rows[0].cells
        hdr[0].text = 'CO'
        hdr[1].text = 'Statement'
        
        for co in course_data['course_outcomes']:
            row = table.add_row().cells
            row[0].text = co['co_id']
            row[1].text = co['statement']
        
        self.doc.add_page_break()
    
    def _add_lesson_plan(self, course_data: dict):
        """Add Course Lesson Plan table"""
        self.doc.add_heading('Course Lesson Plan', level=1)
        
        # Create table with headers
        headers = ['Sr. No.', 'Proposed Date', 'Topics', 'Delivery Mode', 'CO', 'Assessment Tool', 'Ref. book', 'Actual Date']
        
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        
        # Header row
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].runs[0].font.size = Pt(9)
        
        # Data rows
        for item in course_data.get('lesson_plan', []):
            row = table.add_row().cells
            row[0].text = str(item.get('sr_no', ''))
            row[1].text = item.get('proposed_date', '')
            row[2].text = item.get('topic', '')
            row[3].text = item.get('delivery_mode', '')
            row[4].text = ', '.join(item.get('co_mapping', []))
            row[5].text = ', '.join(item.get('assessment_tool', []))
            row[6].text = item.get('reference_book', '')
            row[7].text = item.get('actual_date', '')
        
        self.doc.add_page_break()
    
    def _add_co_po_mapping(self, course_data: dict):
        """Add CO-PO Mapping section"""
        self.doc.add_heading('CO-PO Mapping', level=1)
        
        p = self.doc.add_paragraph()
        p.add_run("(BL – Bloom's Taxonomy, C – Competency, PI – Performance Indicator)")
        
        # Mapping matrix table
        pos = self.global_config['program_outcomes']
        cos = course_data['course_outcomes']
        mappings = {m['co_id']: m for m in course_data.get('co_po_mapping', [])}
        
        # Create matrix table
        table = self.doc.add_table(rows=len(cos) + 1, cols=len(pos) + 1)
        table.style = 'Table Grid'
        
        # Header
        table.rows[0].cells[0].text = 'CO'
        for i, po in enumerate(pos):
            table.rows[0].cells[i + 1].text = po
            table.rows[0].cells[i + 1].paragraphs[0].runs[0].font.size = Pt(8)
        
        # Data
        for i, co in enumerate(cos):
            table.rows[i + 1].cells[0].text = co['co_id']
            mapping = mappings.get(co['co_id'], {}).get('po_mappings', {})
            for j, po in enumerate(pos):
                val = mapping.get(po)
                table.rows[i + 1].cells[j + 1].text = str(val) if val else ''
        
        # Justification table
        self.doc.add_paragraph()
        self.doc.add_heading('CO-PO Mapping Justification', level=2)
        
        just_table = self.doc.add_table(rows=1, cols=5)
        just_table.style = 'Table Grid'
        
        just_headers = ['CO', 'BL', 'PI', 'PO Mapping', 'Justification']
        for i, h in enumerate(just_headers):
            just_table.rows[0].cells[i].text = h
            just_table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        
        for co in course_data.get('co_po_mapping', []):
            row = just_table.add_row().cells
            row[0].text = co.get('co_id', '')
            row[1].text = str(co.get('blooms_level', ''))
            row[2].text = co.get('performance_indicator', '')
            po_map = co.get('po_mappings', {})
            row[3].text = ', '.join([f"{k}:{v}" for k, v in po_map.items() if v])
            row[4].text = co.get('justification', '')
        
        self.doc.add_page_break()
    
    def _add_assessment_tools(self, course_data: dict):
        """Add Assessment Tools section"""
        self.doc.add_heading('Course Assessment', level=1)
        
        assessments = course_data.get('assessments', {})
        
        # ISE1 table
        if 'ise1' in assessments:
            self.doc.add_heading('ISE 1 Assessment', level=2)
            table = self.doc.add_table(rows=1, cols=5)
            table.style = 'Table Grid'
            
            headers = ['Activity', 'Marks', 'Scale To', 'CO Mapped', 'Date']
            for i, h in enumerate(headers):
                table.rows[0].cells[i].text = h
                table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
            
            for item in assessments['ise1']:
                row = table.add_row().cells
                row[0].text = item.get('activity_type', '')
                row[1].text = str(item.get('marks', ''))
                row[2].text = str(item.get('scale_to', ''))
                row[3].text = item.get('co_mapping', '')
                row[4].text = item.get('date', '')
        
        # ISE2 table
        if 'ise2' in assessments:
            self.doc.add_paragraph()
            self.doc.add_heading('ISE 2 Assessment', level=2)
            table = self.doc.add_table(rows=1, cols=5)
            table.style = 'Table Grid'
            
            headers = ['Activity', 'Marks', 'Scale To', 'CO Mapped', 'Date']
            for i, h in enumerate(headers):
                table.rows[0].cells[i].text = h
                table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
            
            for item in assessments['ise2']:
                row = table.add_row().cells
                row[0].text = item.get('activity_type', '')
                row[1].text = str(item.get('marks', ''))
                row[2].text = str(item.get('scale_to', ''))
                row[3].text = item.get('co_mapping', '')
                row[4].text = item.get('date', '')
        
        # Tutorials table
        if 'tutorials' in assessments:
            self.doc.add_paragraph()
            self.doc.add_heading('Tutorial Assessment', level=2)
            table = self.doc.add_table(rows=1, cols=5)
            table.style = 'Table Grid'
            
            headers = ['Tutorial No.', 'Activity', 'Marks', 'Scale To', 'CO']
            for i, h in enumerate(headers):
                table.rows[0].cells[i].text = h
                table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
            
            for item in assessments['tutorials']:
                row = table.add_row().cells
                row[0].text = str(item.get('tutorial_no', ''))
                row[1].text = item.get('activity', '')
                row[2].text = str(item.get('marks', ''))
                row[3].text = str(item.get('scale_to', ''))
                row[4].text = item.get('co_mapping', '')
        
        self.doc.add_page_break()
    
    def _add_recommended_books(self, course_data: dict):
        """Add Recommended Books section"""
        books = course_data.get('recommended_books', [])
        if not books:
            return
        
        self.doc.add_heading('Recommended Books', level=1)
        
        for book in books:
            p = self.doc.add_paragraph()
            p.add_run(f"{book.get('id', '')}. ").bold = True
            p.add_run(f"{book.get('author', '')}: ")
            p.add_run(f'"{book.get("title", "")}", ')
            if book.get('edition'):
                p.add_run(f"{book.get('edition')} Edition, ")
            p.add_run(book.get('publisher', ''))
        
        self.doc.add_page_break()
    
    def _add_exit_survey(self, course_data: dict):
        """Add Course Exit Survey section"""
        survey = course_data.get('exit_survey', {})
        if not survey:
            return
        
        self.doc.add_heading('Course Exit Survey', level=1)
        
        p = self.doc.add_paragraph()
        p.add_run(f"{course_data['course_info']['course_name']} Exit Survey")
        
        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.add_run("To what extent do you agree with the following statements about your learning experience in this course?")
        
        # Questions
        for q in survey.get('questions', []):
            self.doc.add_paragraph()
            p = self.doc.add_paragraph()
            p.add_run(f"{q['co_id']}: ").bold = True
            p.add_run(q['statement'])
            
            # Options
            for option in self.global_config['exit_survey_options']:
                self.doc.add_paragraph(f"○ {option}", style='List Bullet')


def main():
    """Main function to generate course diary"""
    import sys
    
    base_path = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    config_path = os.path.join(base_path, 'config', 'global_config.json')
    
    # Load sample course data
    data_path = os.path.join(base_path, 'data', 'sample_dms_course.json')
    
    with open(data_path, 'r') as f:
        course_data = json.load(f)
    
    # Generate document
    generator = CourseDiaryGenerator(config_path)
    output_path = os.path.join(base_path, 'output', f"Course_Diary_{course_data['course_info']['course_code']}.docx")
    
    # Create output directory
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    generator.create_document(course_data, output_path)
    print(f"Generated: {output_path}")


if __name__ == '__main__':
    main()
